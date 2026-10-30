"""PowerPoint deck and Excel workbook: built from real analyses, opened back and checked (no AI, no network)."""
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))

import make_fixtures  # noqa: E402
import make_s1  # noqa: E402
from analyst import ai, market, pipeline, pptx_report, xlsx_export  # noqa: E402


@pytest.fixture
def no_ai(monkeypatch, tmp_path):
    monkeypatch.setattr(pipeline, "DOCS", tmp_path / "docs")
    monkeypatch.setattr(ai, "active", lambda: None)


def s1(lang="en"):
    return json.loads(json.dumps(pipeline.run({"kind": "upload", "data": make_s1.s1().encode(), "name": "s1.htm", "lang": lang}, lambda *_: None)))


def deck_text(data: bytes) -> tuple[list[str], list]:
    prs = Presentation(io.BytesIO(data))
    texts, charts = [], []
    for slide in prs.slides:
        parts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if sh.has_table:
                parts += [c.text for row in sh.table.rows for c in row.cells]
            if sh.has_chart:
                charts.append(sh.chart)
        texts.append("\n".join(parts))
    return texts, charts


def sheets(data: bytes) -> dict[str, str]:
    """sheet name -> its XML with shared strings inlined (enough to look for labels, numbers and formulas)."""
    z = zipfile.ZipFile(io.BytesIO(data))
    shared = re.findall(r"<t[^>]*>(.*?)</t>", z.read("xl/sharedStrings.xml").decode()) if "xl/sharedStrings.xml" in z.namelist() else []
    names = re.findall(r'<sheet name="([^"]+)"', z.read("xl/workbook.xml").decode())
    out = {}
    for i, name in enumerate(names, start=1):
        xml = z.read(f"xl/worksheets/sheet{i}.xml").decode()
        out[name] = re.sub(r'<c r="[A-Z]+\d+"( s="\d+")? t="s"><v>(\d+)</v>', lambda m: "«" + shared[int(m.group(2))] + "»", xml)
    return out


def test_deck_for_a_prospectus(no_ai):
    data = pptx_report.build(s1(), "ocean")
    texts, charts = deck_text(data)
    assert len(texts) >= 10 and "Example AI, PBC" in texts[0]
    allt = "\n".join(texts)
    for s in ("Headline figures", "$16.3B", "+1,300%", "vs H1 2025", "Key findings", "Financial highlights · Income statement",
              "Arithmetic checks", "How it was read", "p. F-3", "Not financial advice"):
        assert s in allt, s
    # editable charts with the document's numbers (in billions), and the comparative period as a second series
    res = next(c for c in charts if [s.name for s in c.plots[0].series] == ["H1 2026", "H1 2025"])
    cats = list(res.plots[0].categories)
    assert cats[0] == "Revenue" and res.plots[0].series[0].values[0] == pytest.approx(16.29, abs=0.01)
    # every chart slide says where its figures are printed
    assert all("Source:" in tx for tx in texts if "Results" in tx or "Over time" in tx)


def test_deck_in_italian_and_for_a_market_sheet(no_ai):
    texts, _ = deck_text(pptx_report.build(s1("it"), "midnight"))
    assert "I numeri chiave" in texts[1] and any("Fonte:" in tx for tx in texts)
    from synth import series
    a = {**market.build(series()), "lang": "en", "title": "Bitcoin USD", "source": "Yahoo Finance"}
    texts, charts = deck_text(pptx_report.build(json.loads(json.dumps(a)), "neon"))
    assert any(tx.startswith("Price history") for tx in texts) and any(tx.startswith("Return and risk") for tx in texts)
    assert len(charts) >= 2


def test_workbook_numbers_formulas_and_pages(no_ai):
    ws = sheets(xlsx_export.build(s1()))
    assert list(ws) == ["Summary", "Statements", "Ratios", "Checks", "Figures read"]
    st = ws["Statements"]
    assert "«Revenue»" in st and "<v>16286</v>" in st and "<v>1163</v>" in st      # millions, stored as numbers
    assert "«Amounts in USD millions; per share in USD; shares in millions.»" in st
    assert re.search(r'<f>IF\(OR\(B\d+="",B\d+=""\),"",B\d+-ABS\(B\d+\)\)</f><v>3447(\.0)?</v>', st)  # FCF formula
    assert re.search(r"<f>IF\(OR\(C\d+=.*?\(B\d+-C\d+\)/C\d+.*?</f><v>13\.0034", st)                        # change as a formula
    assert "«F-3»" in st and "«H1 2026»" in st and "«Jun 30, 2026»" in st
    assert "«OK»" in ws["Checks"] and "«Assets = liabilities + redeemable preferred + equity»" in ws["Checks"]
    assert "<autoFilter" in ws["Figures read"]


def test_workbook_for_tagged_filing_and_market(no_ai, tmp_path):
    fx = HERE / "fixtures" / "sec_10k.htm"
    if not fx.exists():
        make_fixtures.main() if hasattr(make_fixtures, "main") else None
    a = json.loads(json.dumps(pipeline.run({"kind": "upload", "data": fx.read_bytes(), "name": "x.htm", "lang": "it"}, lambda *_: None)))
    ws = sheets(xlsx_export.build(a))
    assert "Prospetti" in ws and "«Ricavi»" in ws["Prospetti"]
    texts, charts = deck_text(pptx_report.build(a, "forest"))
    assert charts and any("Risultati" in tx for tx in texts)
    from synth import series
    m = json.loads(json.dumps({**market.build(series()), "lang": "it", "title": "BTC"}))
    ws = sheets(xlsx_export.build(m))
    assert "Prezzi" in ws and "Rendimenti" in ws and ws["Prezzi"].count("<row ") > 100


def test_export_routes(no_ai):
    import server
    a = s1()
    client = TestClient(server.app, base_url="http://127.0.0.1:8765")
    r = client.get(f"/api/doc/{a['id']}/report.pptx", params={"theme": "sunset"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    assert 'filename="example-ai-pbc-slides-en.pptx"' in r.headers["content-disposition"] and r.content[:2] == b"PK"
    r = client.get(f"/api/doc/{a['id']}/data.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK" and "spreadsheetml" in r.headers["content-type"]
    assert client.get(f"/api/doc/{a['id']}/report.docx").status_code == 404
    assert client.get("/api/doc/nope/report.pptx").status_code == 404


def test_tagged_quarter_is_named_as_the_issuer_names_it(no_ai):
    """A 10-Q whose fiscal year ends in June: the quarter to March is Q3 FY2025 (dei tags), not "Q1"."""
    html = (HERE / "fixtures" / "sec_10k.htm").read_text()
    html = (html.replace('contextRef="FY25">10-K</ix:nonNumeric>',
                         'contextRef="FY25">10-Q</ix:nonNumeric><ix:nonNumeric name="dei:DocumentFiscalPeriodFocus" contextRef="FY25">Q3</ix:nonNumeric>'
                         '<ix:nonNumeric name="dei:DocumentFiscalYearFocus" contextRef="FY25">2025</ix:nonNumeric>')
            .replace("<xbrli:startDate>2024-07-01</xbrli:startDate><xbrli:endDate>2025-06-30</xbrli:endDate>",
                     "<xbrli:startDate>2025-01-01</xbrli:startDate><xbrli:endDate>2025-03-31</xbrli:endDate>")
            .replace("<xbrli:startDate>2023-07-01</xbrli:startDate><xbrli:endDate>2024-06-30</xbrli:endDate>",
                     "<xbrli:startDate>2024-01-01</xbrli:startDate><xbrli:endDate>2024-03-31</xbrli:endDate>"))
    a = pipeline.run({"kind": "upload", "data": html.encode(), "name": "q.htm", "lang": "en"}, lambda *_: None)
    card = {c["key"]: c for c in a["card"]}
    assert card["revenue"]["vs"] == "Q3 FY2024"
    a = pipeline.run({"kind": "upload", "data": html.encode(), "name": "q.htm", "lang": "it"}, lambda *_: None)
    assert {c["key"]: c for c in a["card"]}["revenue"]["vs"] == "3° trim. es. 2024"
